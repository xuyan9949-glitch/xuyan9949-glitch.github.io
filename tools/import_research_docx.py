#!/usr/bin/env python3
"""Import supplied Word research, preserve content, and synchronize static routes."""
from pathlib import Path
import json, shutil, hashlib, re
from html import escape
from docx import Document
from docx.oxml.ns import qn
from lxml import etree

import argparse
parser=argparse.ArgumentParser(description="Import original company research Word into the existing research library")
parser.add_argument('source',type=Path)
parser.add_argument('--ticker',required=True)
parser.add_argument('--kind',required=True,choices=['dossier','story'])
parser.add_argument('--date',required=True)
parser.add_argument('--source-name')
parser.add_argument('--dry-run',action='store_true',help='Check source and summarize the import without writing files')
args=parser.parse_args()
import datetime
try:
 if datetime.date.fromisoformat(args.date).isoformat()!=args.date:raise ValueError()
except ValueError:
 parser.error('--date must be a valid YYYY-MM-DD date')
if not args.source.is_file():parser.error('Source Word file does not exist')
ROOT=Path(__file__).resolve().parents[1]
assets=ROOT/'images/research'
image_files={}
archive=json.loads((ROOT/'content/archives.json').read_text())
registry=ROOT/'content/research-documents.json'
documents=json.loads(registry.read_text())['documents'] if registry.exists() else []
if not any(c['id']==args.ticker.upper() for c in archive['companies']):
 parser.error('Register company in content/archives.json first')
def inline(element,doc):
 out=''
 for node in element:
  if node.tag==qn('w:r'):
   text=''.join(escape(t.text or '') for t in node.iter(qn('w:t')))
   if node.find(qn('w:br')) is not None:text+='<br>'
   rp=node.find(qn('w:rPr'))
   if rp is not None and rp.find(qn('w:b')) is not None:text='<strong>'+text+'</strong>'
   out+=text
   for blip in node.iter(qn('a:blip')):
    rid=blip.get(qn('r:embed'))
    if rid:
     part=doc.part.related_parts[rid];ext=part.content_type.split('/')[-1].replace('jpeg','jpg')
     name=hashlib.sha256(part.blob).hexdigest()[:20]+'.'+ext
     image_files[name]=part.blob
     out+='<img loading="lazy" src="/images/research/'+name+'" alt="原始研究文稿附图">'
  elif node.tag==qn('w:hyperlink'):
   content=inline(node,doc);rid=node.get(qn('r:id'));url=doc.part.rels[rid].target_ref if rid else ''
   out+=('<a href="'+escape(str(url),quote=True)+'" target="_blank" rel="noopener noreferrer">'+content+'</a>') if str(url).startswith(('https://','http://')) else content
 return out
for ticker in [args.ticker.upper()]:
 c=next(c for c in archive['companies'] if c['id']==ticker)
 c.setdefault('documents',[])
 for kind in [args.kind]:
  files=[args.source]
  source=files[0]
  try:doc=Document(source)
  except Exception as error:parser.error('Cannot read source as a Word DOCX: '+str(error))
  body=[];heading=0
  for element in doc.element.body:
   if element.tag==qn('w:p'):
    text=''.join(element.itertext()) # Heading detection uses plain w:t values below.
    text=''.join(t.text or '' for t in element.iter(qn('w:t'))).strip()
    content=inline(element,doc)
    if not content:continue
    ps=element.find(qn('w:pPr'));style=ps.find(qn('w:pStyle')) if ps is not None else None
    sty=style.get(qn('w:val'),'') if style is not None else ''
    is_head='Heading' in sty or bool(re.fullmatch(r'\d{1,2}\.\s+[^。，；.!?]{1,54}',text)) or text in ['核心摘要','先给核心结论','公司与商业模式','财务与估值','催化剂、风险与市场争议']
    if is_head and not heading:level='h2'
    else:level='h2' if is_head else 'p'
    if is_head:heading+=1
    body.append('<'+level+'>'+content+'</'+level+'>')
   elif element.tag==qn('w:tbl'):
    rows=[]
    for i,row in enumerate(element.findall(qn('w:tr'))):
     celltag='th' if i==0 else 'td';cells=[]
     for cell in row.findall(qn('w:tc')):
      cells.append('<'+celltag+'>'+'<br>'.join(inline(p,doc) for p in cell.findall(qn('w:p')))+' </'+celltag+'>')
     rows.append('<tr>'+''.join(cells)+'</tr>')
    body.append('<div class="v0-table-scroll"><table><tbody>'+''.join(rows)+'</tbody></table></div>')
  original_name=ticker+('_Investment_Dossier.docx' if kind=='dossier' else '_Company_Story.docx')
  dest=ROOT/'downloads/research'/ticker.lower()/original_name
  title=next((p.text for p in doc.paragraphs if p.text.strip() and p.text.strip()!=ticker),ticker)
  item={'id':ticker.lower()+'-'+kind,'ticker':ticker,'kind':kind,'title':title,'date':args.date,'file':'/research/'+ticker.lower()+'/'+kind+'/','download':'/'+dest.relative_to(ROOT).as_posix(),'sourceFile':args.source_name or source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'html':'\n'.join(body),'text':'\n'.join(''.join(t.text or '' for t in e.iter(qn('w:t'))) for e in doc.element.body)}
  if not item['text'].strip():parser.error('Source document contains no readable text')
  if args.dry_run:
   print(json.dumps({'dryRun':True,'action':'replace' if any(d['id']==item['id'] for d in documents) else 'add','id':item['id'],'title':title,'date':args.date,'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'images':len(image_files),'sha256':item['sha256']},ensure_ascii=False,indent=2))
   raise SystemExit(0)
  assets.mkdir(parents=True,exist_ok=True)
  for name,blob in image_files.items():(assets/name).write_bytes(blob)
  dest.parent.mkdir(parents=True,exist_ok=True)
  if source.resolve()!=dest.resolve():shutil.copyfile(source,dest)
  documents=[d for d in documents if d['id']!=item['id']]
  c['documents']=[d for d in c['documents'] if d['id']!=item['id']]
  documents.append(item);c['documents'].append({k:item[k] for k in ['id','kind','title','date','file','download','sourceFile','sha256']})
(ROOT/'content/research-documents.json').write_text(json.dumps({'version':1,'documents':documents},ensure_ascii=False,indent=2)+'\n')
(ROOT/'content/archives.json').write_text(json.dumps(archive,ensure_ascii=False,indent=2)+'\n')
from sync_journal import sync_journal
sync_journal()
print('Imported original research:',args.ticker,args.kind)
