"""Verify import preflight never mutates the destination, even on bad input."""
from pathlib import Path
import json
import hashlib
import shutil
import subprocess
import sys
import tempfile
import unittest
from docx import Document


class ImportPreflight(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        (self.root / 'tools').mkdir()
        (self.root / 'content').mkdir()
        self.script = self.root / 'tools/import_research_docx.py'
        shutil.copyfile(Path(__file__).with_name('import_research_docx.py'), self.script)
        (self.root / 'content/archives.json').write_text(json.dumps({'companies': [{'id': 'MU', 'documents': []}]}))
        (self.root / 'content/research-documents.json').write_text(json.dumps({'version': 1, 'documents': []}))
        self.source = self.root / 'source.docx'
        doc = Document()
        doc.add_heading('新的公司档案', 1)
        doc.add_paragraph('原始观察。')
        doc.add_table(rows=2, cols=2).cell(0, 0).text = '营收'
        doc.save(self.source)

    def tearDown(self):
        self.temp.cleanup()

    def snapshot(self):
        return {str(p.relative_to(self.root)): hashlib.sha256(p.read_bytes()).hexdigest()
                for p in self.root.rglob('*') if p.is_file()}

    def run_import(self, *extra, source=None, date='2026-10-02'):
        before = self.snapshot()
        result = subprocess.run([sys.executable, str(self.script), str(source or self.source),
                                 '--ticker', 'MU', '--kind', 'dossier', '--date', date, *extra],
                                capture_output=True, text=True)
        self.assertEqual(before, self.snapshot(), 'Preflight changed destination files')
        return result

    def test_dry_run_summarizes_without_creating_any_output(self):
        result = self.run_import('--dry-run')
        self.assertEqual(result.returncode, 0, result.stderr)
        info = json.loads(result.stdout)
        self.assertTrue(info['dryRun'])
        self.assertEqual(info['action'], 'add')
        self.assertEqual(info['title'], '新的公司档案')
        self.assertEqual(info['tables'], 1)
        self.assertEqual(info['sha256'], hashlib.sha256(self.source.read_bytes()).hexdigest())

    def test_existing_document_is_reported_as_replacement(self):
        (self.root / 'content/research-documents.json').write_text(json.dumps({'documents': [{'id': 'mu-dossier'}]}))
        result = self.run_import('--dry-run')
        self.assertEqual(json.loads(result.stdout)['action'], 'replace')

    def test_invalid_dates_and_sources_leave_destination_untouched(self):
        for date in ['2026-02-30', '20261002', 'not-a-date']:
            with self.subTest(date=date):
                self.assertNotEqual(self.run_import(date=date).returncode, 0)
        self.assertNotEqual(self.run_import(source=self.root / 'missing.docx').returncode, 0)
        invalid = self.root / 'invalid.docx'
        invalid.write_text('This is not a Word file')
        self.assertNotEqual(self.run_import(source=invalid).returncode, 0)
        empty = self.root / 'empty.docx'
        Document().save(empty)
        self.assertNotEqual(self.run_import(source=empty).returncode, 0)


if __name__ == '__main__':
    unittest.main()
