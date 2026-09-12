#!/usr/bin/env python3
"""Generate genuine offline DOCX/PDF editions of public articles.
Run with the bundled Python runtime; requires python-docx, lxml and LibreOffice.
"""
from pathlib import Path
import argparse,datetime,hashlib,json,re,subprocess,sys,shutil,os,tempfile
from urllib.parse import urljoin,urlparse,unquote
from lxml import html
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image
ROOT=Path(__file__).resolve().parents[1]
GREEN='285B49';GRAY='68736B';FONT='Arial Unicode MS'

def element(tag,**attrs):
 e=OxmlElement(tag)
 for k,v in attrs.items():e.set(qn(k),str(v))
 return e

def add_link(p,text,url):
 link=element('w:hyperlink');link.set(qn('r:id'),p.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True))
 run=element('w:r');props=element('w:rPr');props.append(element('w:color',**{'w:val':GREEN}));props.append(element('w:u',**{'w:val':'single'}));run.append(props);t=element('w:t');t.text=text;run.append(t);link.append(run);p._p.append(link)

def inline(p,node,bold=False,italic=False,base=''):
 if node.text:
  r=p.add_run(node.text);r.bold=bold;r.italic=italic
 for child in node:
  tag=child.tag.lower() if isinstance(child.tag,str) else ''
  if tag=='br':p.add_run().add_break()
  elif tag=='a' and child.get('href'):
   url=urljoin(base,child.get('href'));add_link(p,child.text_content(),url)
  elif tag=='img':add_picture(p,child)
  elif tag not in ['script','style']:inline(p,child,bold or tag in ['b','strong'],italic or tag in ['i','em'],base)
  if child.tail:
   r=p.add_run(child.tail);r.bold=bold;r.italic=italic

def add_picture(p,img):
 src=img.get('src','');url=urlparse(src)
 if url.netloc and url.netloc!='www.xxyalpha.cn':raise ValueError('Remote image requires explicit local asset: '+src)
 path=(ROOT/unquote(url.path.lstrip('/'))).resolve()
 if not path.is_relative_to(ROOT):raise ValueError('Image is outside site')
 if not path.is_file():
  p.add_run('原图暂缺：'+img.get('alt','未命名图片')+'。该图片文件未随原文归档，未在离线版重绘。').italic=True
  print('Missing original image:',src,flush=True)
  return
 with Image.open(path) as im:w,h=im.size
 width=min(16.6,18*w/h)
 p.add_run().add_picture(str(path),width=Cm(width));p.alignment=WD_ALIGN_PARAGRAPH.CENTER
 p.paragraph_format.keep_together=True

def blocks(parent,node,base,depth=0):
 for child in node:
  tag=child.tag.lower() if isinstance(child.tag,str) else ''
  if tag in ['script','style']:continue
  if tag in ['h1','h2','h3','h4','h5','h6']:
   level=min(3,max(1,int(tag[1])-1));p=parent.add_paragraph(style='Heading '+str(level));inline(p,child,base=base)
  elif tag in ['p','blockquote','pre']:
   if not child.text_content().strip() and not child.xpath('.//img'):continue
   if tag=='blockquote' and child.xpath('./p|./ul|./ol'):
    blocks(parent,child,base,depth+1);continue
   p=parent.add_paragraph();inline(p,child,base=base)
   if depth or tag=='blockquote':p.paragraph_format.left_indent=Cm(.4)
   if tag=='pre':
    p.paragraph_format.space_after=Pt(9);p.paragraph_format.line_spacing=Pt(12)
    for r in p.runs:r.font.size=Pt(8);r.font.name='Menlo'
  elif tag in ['ul','ol']:
   for i,li in enumerate(child.findall('li')):
    p=parent.add_paragraph();p.paragraph_format.left_indent=Cm(.45);p.paragraph_format.first_line_indent=Cm(-.35);p.add_run(str(i+1)+'. ' if tag=='ol' else '• ')
    shallow=html.fromstring(html.tostring(li,encoding='unicode'))
    nested=shallow.xpath('./ul|./ol')
    for n in nested:shallow.remove(n)
    inline(p,shallow,base=base)
    for n in li.xpath('./ul|./ol'):
     wrapper=html.Element('div');wrapper.append(html.fromstring(html.tostring(n)));blocks(parent,wrapper,base,depth+1)
  elif tag=='table':
   rows=child.xpath('./tr|./thead/tr|./tbody/tr|./tfoot/tr')
   if not rows:continue
   occupied=set();placements=[];cols=0
   for ri,row in enumerate(rows):
    ci=0
    for source in row.xpath('./th|./td'):
     while (ri,ci) in occupied:ci+=1
     rs=int(source.get('rowspan','1'));cs=int(source.get('colspan','1'))
     placements.append((ri,ci,rs,cs,source))
     for rr in range(ri,ri+rs):
      for cc in range(ci,ci+cs):occupied.add((rr,cc))
     ci+=cs;cols=max(cols,ci)
   table=parent.add_table(rows=len(rows),cols=cols);table.style='Table Grid';table.autofit=False
   for col in table.columns:col.width=Cm(16.6/cols)
   for ri,ci,rs,cs,source in placements:
    cell=table.cell(ri,ci)
    if rs>1 or cs>1:cell=cell.merge(table.cell(ri+rs-1,ci+cs-1))
    p=cell.paragraphs[0];inline(p,source,base=base);p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4);p.paragraph_format.line_spacing=Pt(14)
    for run in p.runs:run.font.size=Pt(9 if cols>4 else 10);run.bold=source.tag=='th' or run.bold
    if source.tag=='th':cell._tc.get_or_add_tcPr().append(element('w:shd',**{'w:fill':'E9EEE6'}))
   if rows[0].xpath('./th'):table.rows[0]._tr.get_or_add_trPr().append(element('w:tblHeader'))
   parent.add_paragraph().paragraph_format.space_after=Pt(3)
  elif tag=='img':add_picture(parent.add_paragraph(),child)
  elif tag=='hr':continue
  else:
   if child.text and child.text.strip():parent.add_paragraph(child.text.strip())
   blocks(parent,child,base,depth)
  if child.tail and child.tail.strip():parent.add_paragraph(child.tail.strip())

def make_docx(article,out):
 src=(ROOT/article['file'].strip('/')/'index.html').read_text();tree=html.fromstring(src);body=tree.xpath('//*[contains(concat(" ",normalize-space(@class)," ")," article-content ")]')[0]
 doc=Document();sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7);sec.top_margin=Cm(1.9);sec.bottom_margin=Cm(1.8);sec.left_margin=Cm(2.2);sec.right_margin=Cm(2.2)

 for style in doc.styles:
  if style.type in [1,2]:
   style.font.name=FONT
   rp=style._element.get_or_add_rPr();rf=rp.rFonts
   if rf is not None:
    for attr in list(rf.attrib):
     if 'Theme' in attr or 'theme' in attr:del rf.attrib[attr]
    rf.set(qn('w:eastAsia'),FONT)
   for border in style._element.xpath('.//w:pBdr'):border.getparent().remove(border)
 doc.styles['Caption'].font.color.rgb=RGBColor.from_string(GRAY)
 doc.styles['Caption'].font.bold=False
 doc.styles['Caption'].paragraph_format.line_spacing=Pt(14)
 compact=article['id'] in {'nvidia-25b-bond-issuance','daodejing-chapter-4','aaoi-insider-selling-analysis','ccl-ai-material-2026','semi-analysis-ai-gpu-scarcity','avgo-vs-mrvl-platform-vs-chip','gtc-taipei-2026'}
 normal=doc.styles['Normal'];normal.font.name=FONT;normal.font.size=Pt(11);normal._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);normal.paragraph_format.line_spacing=Pt(17 if article['id']=='nvidia-25b-bond-issuance' else 18 if article['id'] in {'nvidia-25b-bond-issuance','daodejing-chapter-4','aaoi-insider-selling-analysis','ccl-ai-material-2026','semi-analysis-ai-gpu-scarcity','avgo-vs-mrvl-platform-vs-chip','gtc-taipei-2026'} else 19);normal.paragraph_format.space_after=Pt(5 if compact else 8);normal.paragraph_format.widow_control=True
 for name,size in [('Title',22),('Heading 1',16),('Heading 2',13),('Heading 3',12)]:
  style=doc.styles[name];style.font.name=FONT;style.font.size=Pt(size);style.font.color.rgb=RGBColor.from_string('000000' if name=='Title' else GREEN);style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),FONT);style.paragraph_format.space_before=Pt(14 if name!='Title' else 0);style.paragraph_format.space_after=Pt(8);style.paragraph_format.keep_with_next=True
 header=sec.header.paragraphs[0];header.text='XXY Alpha  /  Leslie 的研究笔记';header.style='Caption';header.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
 footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;r=footer.add_run('XXY Alpha  ·  ');r.font.size=Pt(8);field=element('w:fldSimple',**{'w:instr':'PAGE'});footer._p.append(field)
 doc.add_paragraph(article['title'],style='Title')
 stamp=datetime.date.today().isoformat();p=doc.add_paragraph(f"原文日期 {article['date']}    文件生成 {stamp}");p.style='Caption'
 base='https://www.xxyalpha.cn'+article['file'];p=doc.add_paragraph();add_link(p,'阅读网站原文',base)
 p=doc.add_paragraph('离线阅读版保留原文观点与日期，不代表当前判断。来源链接可点击；网页的最新修订以原文为准。');p.style='Caption'
 blocks(doc,body,base)
 links=[]
 for a in body.xpath('.//a[@href]'):
  u=urljoin(base,a.get('href'))
  if u.startswith(('https://','http://')) and u not in links:links.append(u)
 if links:
  doc.add_paragraph('原文引用链接',style='Heading 1')
  for i,u in enumerate(links):
   p=doc.add_paragraph();p.paragraph_format.line_spacing=Pt(13);p.add_run(str(i+1)+'. ');add_link(p,u,u)
   for run in p.runs:run.font.size=Pt(9)
 doc.core_properties.title=article['title'];doc.core_properties.author='XXY Alpha';doc.core_properties.subject='历史研究文章离线阅读版';doc.save(out)
 return hashlib.sha256(re.search(r'<main\b[^>]*>(.*?)</main>',src,re.S).group(1).encode()).hexdigest()

def configure_fonts(qa_dir):
 if sys.platform!='darwin':return
 fonts=Path(qa_dir)/'_fonts';fonts.mkdir(parents=True,exist_ok=True)
 source=Path('/System/Library/Fonts/Supplemental/Arial Unicode.ttf')
 if not source.is_file():raise FileNotFoundError('A Chinese-capable font is required for export')
 shutil.copyfile(source,fonts/source.name)
 conf=fonts/'fonts.conf'
 conf.write_text('<?xml version="1.0"?><fontconfig><dir>'+str(fonts)+'</dir><cachedir>'+str(fonts/'cache')+'</cachedir><alias><family>sans-serif</family><prefer><family>Arial Unicode MS</family></prefer></alias></fontconfig>')
 os.environ['FONTCONFIG_FILE']=str(conf)

def main():
 ap=argparse.ArgumentParser();ap.add_argument('--id');ap.add_argument('--render',action='store_true');ap.add_argument('--qa-dir',default='/private/tmp/xxyalpha-exports-qa');ap.add_argument('--renderer',default='/Users/xuyan/.codex/plugins/cache/openai-primary-runtime/documents/26.909.22227/skills/documents/render_docx.py');args=ap.parse_args();configure_fonts(args.qa_dir)
 articles=json.loads(re.search(r'const articles = (\[[\s\S]*?\]);',(ROOT/'js/articles.js').read_text()).group(1));manifest_path=ROOT/'content/downloads.json';manifest=json.loads(manifest_path.read_text()) if manifest_path.exists() else {}
 for a in articles:
  if args.id and a['id']!=args.id:continue
  target=ROOT/'downloads'/a['id'];target.mkdir(parents=True,exist_ok=True);docx=target/'article.docx';hash_=make_docx(a,docx)
  if args.render:
   qa=Path(args.qa_dir)/a['id'];subprocess.run([sys.executable,args.renderer,str(docx),'--output_dir',str(qa),'--emit_pdf','--dpi','95'],check=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT)
   shutil.copyfile(qa/'article.pdf',target/'article.pdf')
   manifest[a['id']]={'sourceHash':hash_,'pdf':'/downloads/'+a['id']+'/article.pdf','docx':'/downloads/'+a['id']+'/article.docx','generatedAt':datetime.date.today().isoformat()}
  if args.render:
   manifest_path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n');(ROOT/'js/downloads.js').write_text('window.articleDownloads = '+json.dumps(manifest,ensure_ascii=False,indent=2)+';\n')
  print(a['id'], 'DOCX + PDF' if args.render else 'DOCX',flush=True)
 if args.render:
  manifest_path.write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n');(ROOT/'js/downloads.js').write_text('window.articleDownloads = '+json.dumps(manifest,ensure_ascii=False,indent=2)+';\n')
if __name__=='__main__':main()
